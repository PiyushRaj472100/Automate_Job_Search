"""Comprehensive tests for resume ingestion, parsing, extraction, versioning, and API endpoints."""

import io
import uuid

import pytest
from docx import Document
from fastapi.testclient import TestClient
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from sqlalchemy.orm import Session

from backend.ingestion.extractor import GeminiResumeExtractor, ResumeExtractionError
from backend.ingestion.parser import (
    CorruptFileError,
    EmptyResumeError,
    UnsupportedFileFormatError,
    parse_resume_bytes,
)
from backend.ingestion.schemas import StructuredResumeExtraction
from backend.ingestion.service import ResumeIngestionService


# Helper functions to build real binary test files
def build_sample_pdf(name: str, education: str, skills_str: str) -> bytes:
    """Generate a real, valid PDF in memory using ReportLab."""
    buffer = io.BytesIO()
    c = canvas.Canvas(buffer, pagesize=letter)
    c.drawString(100, 750, f"Candidate Name: {name}")
    c.drawString(100, 725, "Email: candidate@example.com | Phone: +91 9876543210")
    c.drawString(100, 700, f"Education: {education}")
    c.drawString(100, 675, f"Technical Skills: {skills_str}")
    c.drawString(100, 650, "Projects: E-Commerce Microservices in Python, FastAPI, and PostgreSQL.")
    c.drawString(100, 625, "Experience: Software Engineering Intern at TechLabs (6 months).")
    c.save()
    return buffer.getvalue()


def build_sample_docx(name: str, education: str, skills_str: str) -> bytes:
    """Generate a real, valid DOCX in memory using python-docx."""
    doc = Document()
    doc.add_heading(name, level=1)
    doc.add_paragraph("Email: candidate@example.com | Bengaluru, India")
    doc.add_heading("Education", level=2)
    doc.add_paragraph(education)
    doc.add_heading("Technical Skills", level=2)
    doc.add_paragraph(skills_str)
    doc.add_heading("Projects", level=2)
    doc.add_paragraph("Personal Job Portal: Built using Python, FastAPI, and Docker.")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# 1. PDF Parsing Test
def test_pdf_parsing():
    """Verify that a real PDF document is parsed, text cleaned, and SHA-256 computed."""
    pdf_bytes = build_sample_pdf(
        name="Aarav Sharma",
        education="Bachelor of Technology in Computer Science, 2024",
        skills_str="Python, FastAPI, PostgreSQL, Docker, AWS",
    )
    result = parse_resume_bytes("aarav_sharma.pdf", pdf_bytes)

    assert result.file_name == "aarav_sharma.pdf"
    assert result.file_extension == "pdf"
    assert len(result.file_hash) == 64  # SHA-256
    assert "Aarav Sharma" in result.cleaned_text
    assert "FastAPI" in result.cleaned_text
    assert "2024" in result.cleaned_text


# 2. DOCX Parsing Test
def test_docx_parsing():
    """Verify that a real DOCX document is parsed, text cleaned, and SHA-256 computed."""
    docx_bytes = build_sample_docx(
        name="Priya Patel",
        education="Bachelor of Engineering in Information Technology, 2023",
        skills_str="TypeScript, React, Node.js, MongoDB, Docker",
    )
    result = parse_resume_bytes("priya_patel.docx", docx_bytes)

    assert result.file_name == "priya_patel.docx"
    assert result.file_extension == "docx"
    assert len(result.file_hash) == 64
    assert "Priya Patel" in result.cleaned_text
    assert "TypeScript" in result.cleaned_text


# 3. Empty Resume Handling
def test_empty_resume_handling():
    """Verify that empty files or files without extractable text raise EmptyResumeError."""
    # 0-byte file
    with pytest.raises(EmptyResumeError, match="empty"):
        parse_resume_bytes("empty.pdf", b"")

    # Blank PDF with no text
    buffer = io.BytesIO()
    c = canvas.Canvas(buffer, pagesize=letter)
    c.showPage()
    c.save()
    blank_pdf = buffer.getvalue()

    with pytest.raises(EmptyResumeError, match="No extractable text"):
        parse_resume_bytes("blank.pdf", blank_pdf)


# 4. Malformed and Unsupported Files
def test_malformed_and_unsupported_files():
    """Verify corrupt files and unsupported formats are rejected with explicit errors."""
    corrupt_bytes = b"NOT_A_REAL_PDF_HEADER_JUST_GARBAGE"

    # Corrupt PDF
    with pytest.raises(CorruptFileError):
        parse_resume_bytes("corrupt.pdf", corrupt_bytes)

    # Corrupt DOCX
    with pytest.raises(CorruptFileError):
        parse_resume_bytes("corrupt.docx", corrupt_bytes)

    # Unsupported format
    with pytest.raises(UnsupportedFileFormatError, match="Unsupported file format"):
        parse_resume_bytes("resume.txt", b"Plain text resume")


# 5. Structured Gemini Output Validation
def test_structured_gemini_output_validation():
    """Verify strict Pydantic validation of structured JSON output from Gemini."""
    valid_gemini_payload = {
        "education": [
            {
                "degree": "B.Tech Computer Science",
                "institution": "National Institute of Technology",
                "graduation_year": 2024,
                "field_of_study": "Computer Science",
            }
        ],
        "degree": "B.Tech Computer Science",
        "graduation_year": 2024,
        "experience": [],
        "internships": [
            {
                "title": "Backend Engineering Intern",
                "organization": "CloudScale",
                "duration": "6 months",
                "bullets": ["Implemented REST APIs with FastAPI"],
            }
        ],
        "skills": {
            "programming_languages": ["Python", "SQL"],
            "frameworks": ["FastAPI"],
            "libraries": ["Pydantic", "SQLAlchemy"],
            "databases": ["PostgreSQL"],
            "cloud": ["AWS"],
            "ai_ml": [],
            "genai_llm": [],
            "backend": ["FastAPI"],
            "frontend": [],
            "devops": ["Docker"],
            "tools": ["Git"],
            "apis": ["REST"],
        },
        "certifications": ["AWS Certified Cloud Practitioner"],
        "projects": [
            {
                "title": "Job Finder Engine",
                "description": "API service for job intelligence",
                "technologies": ["Python", "FastAPI", "PostgreSQL"],
            }
        ],
        "project_technologies": ["Python", "FastAPI", "PostgreSQL"],
        "domains": ["Developer Tools", "Cloud Computing"],
        "likely_target_roles": ["Junior Backend Engineer", "Python Developer (0-2 YOE)"],
        "seniority": "Entry-Level (0-2 years)",
    }

    validated = GeminiResumeExtractor.validate_extraction_json(valid_gemini_payload)
    assert isinstance(validated, StructuredResumeExtraction)
    assert validated.degree == "B.Tech Computer Science"
    assert validated.graduation_year == 2024
    assert "python" in validated.all_unique_skills()
    assert "fastapi" in validated.all_unique_skills()

    # Test malformed payload
    invalid_payload = {"education": "Not a list"}  # violates schema
    with pytest.raises(ResumeExtractionError, match="Schema validation failed"):
        GeminiResumeExtractor.validate_extraction_json(invalid_payload)


# 6. Missing Fields and Zero Hallucination
def test_missing_fields_handling():
    """Verify that when a resume lacks certifications or internships, they default to empty lists."""
    minimal_payload = {
        "degree": "B.S. in Software Engineering",
        "skills": {
            "programming_languages": ["Python"],
        },
    }
    validated = GeminiResumeExtractor.validate_extraction_json(minimal_payload)
    assert validated.education == []
    assert validated.internships == []
    assert validated.certifications == []
    assert validated.skills.cloud == []
    assert validated.skills.programming_languages == ["Python"]


# 7. Persistence & Duplicate Upload Handling
def test_ingestion_service_persistence_and_versioning(db_session: Session):
    """Verify database persistence of Resume, ResumeProfile, and skills, with duplicate versioning."""
    pdf_bytes = build_sample_pdf(
        name="Vikram Rao",
        education="Bachelor of Technology in Computer Science, 2024",
        skills_str="Python, FastAPI, PostgreSQL, Docker",
    )

    service = ResumeIngestionService()

    # 1. First upload: Creates new resume and profile
    resume1, profile1, data1, is_dup1 = service.process_and_persist_resume(
        db=db_session,
        file_name="vikram_rao.pdf",
        content=pdf_bytes,
        profile_name="Vikram Primary Profile",
    )

    assert is_dup1 is False
    assert isinstance(resume1.id, uuid.UUID)
    assert isinstance(profile1.id, uuid.UUID)
    assert profile1.resume_id == resume1.id
    assert len(profile1.skills) >= 3

    # 2. Second upload of identical file bytes: Reuses resume record and creates versioned profile
    resume2, profile2, data2, is_dup2 = service.process_and_persist_resume(
        db=db_session,
        file_name="vikram_rao.pdf",
        content=pdf_bytes,
    )

    assert is_dup2 is True
    assert resume2.id == resume1.id  # Same resume entity reused
    assert profile2.id != profile1.id  # New distinct profile created
    assert "v2" in profile2.profile_name


# 8. Complete API Endpoints Verification
def test_resume_api_endpoints_e2e(client: TestClient):
    """Verify POST /resumes, GET /resumes, GET /resumes/{id}, and GET /resumes/{id}/profile."""
    pdf_bytes = build_sample_pdf(
        name="Ananya Sen",
        education="Bachelor of Science in Computer Science, 2024",
        skills_str="Python, FastAPI, PostgreSQL, Docker, Git",
    )

    # 1. POST /resumes
    files = {"file": ("ananya_sen.pdf", pdf_bytes, "application/pdf")}
    data = {"profile_name": "Ananya Backend Profile"}
    response = client.post("/resumes", files=files, data=data)
    assert response.status_code == 201
    res_json = response.json()
    assert "resume_id" in res_json
    assert "profile_id" in res_json
    assert res_json["file_name"] == "ananya_sen.pdf"
    assert res_json["is_duplicate"] is False

    resume_id = res_json["resume_id"]

    # 2. GET /resumes
    list_response = client.get("/resumes")
    assert list_response.status_code == 200
    resumes_list = list_response.json()
    assert any(r["id"] == resume_id for r in resumes_list)

    # 3. GET /resumes/{id}
    detail_response = client.get(f"/resumes/{resume_id}")
    assert detail_response.status_code == 200
    detail_json = detail_response.json()
    assert detail_json["file_name"] == "ananya_sen.pdf"
    assert len(detail_json["profiles"]) >= 1

    # 4. GET /resumes/{id}/profile
    profile_response = client.get(f"/resumes/{resume_id}/profile")
    assert profile_response.status_code == 200
    profile_json = profile_response.json()
    assert profile_json["resume_id"] == resume_id
    assert "fastapi" in [s.lower() for s in profile_json["skills"]]
    assert profile_json["structured_data"] is not None

    # 5. GET non-existent resume
    non_existent = str(uuid.uuid4())
    not_found_res = client.get(f"/resumes/{non_existent}")
    assert not_found_res.status_code == 404
